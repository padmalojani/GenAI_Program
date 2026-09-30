import os
import re
from datetime import datetime
from pathlib import Path

try:
    import pyautogui
except KeyError as error:
    if error.args == ("DISPLAY",):
        raise SystemExit(
            "PyAutoGUI needs an active graphical desktop. DISPLAY is unset; "
            "run this script from a desktop session or configure GUI forwarding."
        ) from error
    raise

try:
    import pyperclip
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError as error:
    raise SystemExit(
        "Install the required packages with "
        "'.\\venv\\Scripts\\python.exe -m pip install -r requirements.txt'."
    ) from error

SOURCE_URL = "https://www.goodreturns.in/gold-rates/chennai.html"
REPORT_DIR = Path(__file__).resolve().parent


def copy_gold_rates_from_chrome():
    pyautogui.hotkey("win", "r")
    pyautogui.write("chrome\n", interval=0.05)
    pyautogui.sleep(3)
    pyautogui.hotkey("ctrl", "t")
    pyautogui.write(SOURCE_URL, interval=0.02)
    pyautogui.press("enter")
    pyautogui.sleep(8)

    pyperclip.copy("")
    pyautogui.hotkey("ctrl", "a")
    pyautogui.hotkey("ctrl", "c")
    pyautogui.sleep(1)
    copied_text = pyperclip.paste()
    if not copied_text.strip():
        raise SystemExit("No page text was copied. Check Chrome and try again.")

    lines = [line.strip() for line in copied_text.splitlines() if line.strip()]
    rate_pattern = re.compile(
        r"gold|22\s*(?:k|carat|karat)|24\s*(?:k|carat|karat)|₹|INR|Rs\.?\s*[\d,]",
        re.IGNORECASE,
    )
    rate_lines = [line for line in lines if rate_pattern.search(line)]
    if rate_lines:
        return "\n".join(rate_lines[:30])[:4000]
    return "\n".join(lines[:40])[:4000]


def save_report(timestamp, gold_rates):
    report_path = REPORT_DIR / f"daily_report_{timestamp:%Y-%m-%d}.xlsx"
    if report_path.exists():
        workbook = load_workbook(report_path)
        worksheet = workbook["Daily Report"] if "Daily Report" in workbook.sheetnames else workbook.active
        worksheet.title = "Daily Report"
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Daily Report"

    if worksheet.max_row == 1 and worksheet["A1"].value is None:
        worksheet.append(["Date & Time", "Fetched Gold Rates (Chennai)", "Comment"])
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="287C72")

    worksheet.append([timestamp, gold_rates, "Review rates before making a purchase."])
    row_number = worksheet.max_row
    worksheet.cell(row=row_number, column=1).number_format = "yyyy-mm-dd hh:mm:ss"
    for cell in worksheet[row_number]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    worksheet.column_dimensions["A"].width = 22
    worksheet.column_dimensions["B"].width = 88
    worksheet.column_dimensions["C"].width = 42
    worksheet.row_dimensions[row_number].height = 180
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    workbook.save(report_path)
    return report_path


def save_excel_screenshot(report_path, timestamp):
    os.startfile(str(report_path))
    pyautogui.sleep(6)
    pyautogui.hotkey("win", "up")
    pyautogui.hotkey("ctrl", "home")
    pyautogui.sleep(1)
    screenshot_path = REPORT_DIR / f"daily_report_{timestamp:%Y-%m-%d}.png"
    pyautogui.screenshot().save(str(screenshot_path))
    return screenshot_path


def main():
    pyautogui.FAILSAFE = True
    pyautogui.PAUSE = 0.5
    timestamp = datetime.now()
    print(f"Reading Chennai gold rates at {timestamp:%Y-%m-%d %H:%M:%S}")
    gold_rates = copy_gold_rates_from_chrome()
    report_path = save_report(timestamp, gold_rates)
    print(f"Saved Excel report: {report_path}")
    screenshot_path = save_excel_screenshot(report_path, timestamp)
    print(f"Saved Excel screenshot: {screenshot_path}")


if __name__ == "__main__":
    main()
